import os
import smtplib
from email.message import EmailMessage

from docx import Document
from openpyxl import Workbook, load_workbook

ARCHIVO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Inventario.xlsx")

CORREO_USUARIO = "cristianhernandez.ipedep@gmail.com"
CORREO_CLAVE = ""
CORREO_COORDINADOR = "erickorellana412@gmail.com"

COLUMNAS = {
    "productos": ["nombre", "precio", "existencia"],
    "clientes": ["nombre", "nit", "direccion"],
    "pedidos": ["nombre cliente", "nombre producto", "cantidad producto", "valor pedido"],
}


def crear_libro():
    """Crea Inventario.xlsx con 3 hojas si todavía no existe."""
    if os.path.exists(ARCHIVO):
        return
    libro = Workbook()
    libro.remove(libro.active)
    for hoja, columnas in COLUMNAS.items():
        libro.create_sheet(hoja).append(columnas)
    libro.save(ARCHIVO)


def leer(hoja):
    """Devuelve las filas de una hoja (sin el encabezado)."""
    libro = load_workbook(ARCHIVO)
    filas = []
    for fila in libro[hoja].iter_rows(min_row=2, values_only=True):
        if fila[0] is not None:
            filas.append(list(fila))
    libro.close()
    return filas


def guardar(hoja, filas):
    """Reemplaza todas las filas de una hoja por las que recibe."""
    libro = load_workbook(ARCHIVO)
    hoja_excel = libro[hoja]
    hoja_excel.delete_rows(2, hoja_excel.max_row)
    for fila in filas:
        hoja_excel.append(fila)
    libro.save(ARCHIVO)
    libro.close()


def enviar_correo(destino, asunto, cuerpo, archivo_adjunto):
    """Envía un correo con un archivo adjunto usando Gmail."""
    mensaje = EmailMessage()
    mensaje["From"] = CORREO_USUARIO
    mensaje["To"] = destino
    mensaje["Subject"] = asunto
    mensaje.set_content(cuerpo)

    with open(archivo_adjunto, "rb") as f:
        mensaje.add_attachment(
            f.read(),
            maintype="application",
            subtype="octet-stream",
            filename=os.path.basename(archivo_adjunto),
        )

    with smtplib.SMTP("smtp.gmail.com", 587, timeout=30) as servidor:
        servidor.starttls()
        servidor.login(CORREO_USUARIO, CORREO_CLAVE)
        servidor.send_message(mensaje)


def crear_cotizacion(cliente, producto, precio):
    """Crea cotización.docx reemplazando los datos entre arrobas."""
    lineas = [
        "Estimado @Nombre del cliente@",
        "El precio de nuestro producto @Nombre del producto@",
        "es @Precio del producto@",
    ]
    documento = Document()
    for linea in lineas:
        linea = linea.replace("@Nombre del cliente@", cliente)
        linea = linea.replace("@Nombre del producto@", producto)
        linea = linea.replace("@Precio del producto@", "Q" + format(precio, ",.2f"))
        documento.add_paragraph(linea)

    ruta = os.path.join(os.path.dirname(ARCHIVO), "cotización.docx")
    documento.save(ruta)
    return ruta