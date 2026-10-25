import os
import tempfile
import unittest
from unittest import mock

from docx import Document
from openpyxl import load_workbook

import datos


class PruebasDatos(unittest.TestCase):
    def setUp(self):
        
        self.ruta = os.path.join(self.carpeta.name, "Inventario.xlsx")
        self.parche = mock.patch.object(datos, "ARCHIVO", self.ruta)
        self.parche.start()

    def tearDown(self):
        self.parche.stop()
        self.carpeta.cleanup()

    # Excel
    def test_crear_libro_crea_tres_hojas(self):
        datos.crear_libro()
        libro = load_workbook(self.ruta)
        self.assertEqual(libro.sheetnames, ["productos", "clientes", "pedidos"])
        libro.close()

    def test_crear_libro_pone_encabezados(self):
        datos.crear_libro()
        libro = load_workbook(self.ruta)
        for hoja, columnas in datos.COLUMNAS.items():
            encabezado = [celda.value for celda in libro[hoja][1]]
            self.assertEqual(encabezado, columnas)
        libro.close()

    def test_crear_libro_no_borra_datos_existentes(self):
        datos.crear_libro()
        datos.guardar("productos", [["Mouse", 25.5, 10]])
        datos.crear_libro()  # segunda vez: no debe borrar nada
        self.assertEqual(datos.leer("productos"), [["Mouse", 25.5, 10]])

    def test_hojas_nuevas_estan_vacias(self):
        datos.crear_libro()
        for hoja in datos.COLUMNAS:
            self.assertEqual(datos.leer(hoja), [])

    def test_guardar_y_leer_productos(self):
        datos.crear_libro()
        datos.guardar("productos", [["Mouse", 25.5, 10], ["Teclado", 100.0, 3]])
        self.assertEqual(datos.leer("productos"), [["Mouse", 25.5, 10], ["Teclado", 100.0, 3]])

    def test_guardar_reemplaza_lo_anterior(self):
        datos.crear_libro()
        datos.guardar("productos", [["A", 1.0, 1], ["B", 2.0, 2]])
        datos.guardar("productos", [["C", 3.0, 3]])
        self.assertEqual(datos.leer("productos"), [["C", 3.0, 3]])

    def test_guardar_lista_vacia_deja_la_hoja_vacia(self):
        datos.crear_libro()
        datos.guardar("productos", [["A", 1.0, 1]])
        datos.guardar("productos", [])
        self.assertEqual(datos.leer("productos"), [])

    def test_guardar_y_leer_clientes_y_pedidos(self):
        datos.crear_libro()
        datos.guardar("clientes", [["Ana López", "1234567", "Jalapa"]])
        datos.guardar("pedidos", [["Ana López", "Mouse", 3, 76.5]])
        self.assertEqual(datos.leer("clientes"), [["Ana López", "1234567", "Jalapa"]])
        self.assertEqual(datos.leer("pedidos"), [["Ana López", "Mouse", 3, 76.5]])

    # Word
    def test_cotizacion_tiene_los_datos_correctos(self):
        ruta = datos.crear_cotizacion("Ana López", "Mouse", 25.5)
        lineas = [parrafo.text for parrafo in Document(ruta).paragraphs]
        self.assertEqual(lineas, [
            "Estimado Ana López",
            "El precio de nuestro producto Mouse",
            "es Q25.50",
        ])

    def test_cotizacion_no_deja_arrobas(self):
        ruta = datos.crear_cotizacion("Luis", "Monitor", 1200)
        texto = " ".join(parrafo.text for parrafo in Document(ruta).paragraphs)
        self.assertNotIn("@", texto)

    def test_cotizacion_formatea_miles(self):
        ruta = datos.crear_cotizacion("Luis", "Laptop", 4500)
        lineas = [parrafo.text for parrafo in Document(ruta).paragraphs]
        self.assertEqual(lineas[2], "es Q4,500.00")

    #  Correo
    def test_enviar_correo_con_adjunto(self):
        adjunto = os.path.join(self.carpeta.name, "archivo.txt")
        with open(adjunto, "w") as f:
            f.write("contenido")

        
        with mock.patch("datos.smtplib.SMTP") as smtp_falso:
            datos.enviar_correo("cliente@ejemplo.com", "Asunto de prueba", "Cuerpo", adjunto)

        servidor = smtp_falso.return_value.__enter__.return_value
        servidor.starttls.assert_called_once()
        servidor.login.assert_called_once_with(datos.CORREO_USUARIO, datos.CORREO_CLAVE)
        mensaje = servidor.send_message.call_args[0][0]
        self.assertEqual(mensaje["To"], "cliente@ejemplo.com")
        self.assertEqual(mensaje["Subject"], "Asunto de prueba")
        nombres = [parte.get_filename() for parte in mensaje.iter_attachments()]
        self.assertEqual(nombres, ["archivo.txt"])

    def test_enviar_correo_falla_si_no_hay_conexion(self):
        adjunto = os.path.join(self.carpeta.name, "archivo.txt")
        with open(adjunto, "w") as f:
            f.write("contenido")
        with mock.patch("datos.smtplib.SMTP", side_effect=OSError("sin internet")):
            with self.assertRaises(OSError):
                datos.enviar_correo("a@b.com", "x", "y", adjunto)

    def test_enviar_correo_falla_si_el_adjunto_no_existe(self):
        with mock.patch("datos.smtplib.SMTP"):
            with self.assertRaises(FileNotFoundError):
                datos.enviar_correo("a@b.com", "x", "y", "no_existe.xlsx")


if __name__ == "__main__":
    unittest.main()